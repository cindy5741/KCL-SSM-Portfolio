import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font

columns = ['編號', '擬真問句', '標準問句']

df1 = pd.read_excel('../preprocess_data/AI擬真問句_Q1-Q100_250602.xlsx', header=None)
df2 = pd.read_excel('../preprocess_data/AI擬真問句_Q101-Q135_250609.xlsx', header=None)
df3 = pd.read_excel('../preprocess_data/AI擬真問句_Q136-Q250_250616.xlsx', header=None)

df1.columns = columns
df2.columns = columns
df3.columns = columns
print(f"Number of records in df1 before merging: {len(df1)}")
print(f"Number of records in df2 before merging: {len(df2)}")
print(f"Number of records in df3 before merging: {len(df3)}")
df_merged = pd.concat([df1, df2, df3])

print(f"total number of records: {len(df_merged)}")
# save concat data to Excel
df_merged.to_excel('../preprocess_data/merged_original_data.xlsx', index=False)

wb = load_workbook('../preprocess_data/merged_original_data.xlsx')
ws = wb.active

# set font with size 12
font = Font(name='新細明體', size=12)

# Apply the font to all cells
for row in ws.iter_rows():
    for cell in row:
        cell.font = font

# Save as an Excel file with applied fonts
wb.save('../preprocess_data/merged_original_data.xlsx')

print("Merging completed. Results have been saved to merged_original_data.xlsx")


